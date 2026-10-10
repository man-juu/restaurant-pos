"""Read and write CSV/XLSX tables safely (FR-IMP-001, FR-IMP-003).

Reading: size, row and column limits; XLSX is a zip, so its uncompressed size is checked before
openpyxl opens it (zip-bomb guard); values only, formulas are never evaluated.
Writing: cells that start with = + - @ (or tab/CR) get a leading apostrophe so a spreadsheet
never runs them as formulas (CSV/formula injection)."""

import csv
import io
import zipfile
from collections.abc import Iterable, Sequence
from typing import Any

from openpyxl import Workbook, load_workbook

from app.core.errors import AppError

MAX_ROWS = 2000
MAX_COLUMNS = 40
MAX_CELL = 2000
MAX_UNZIPPED = 50 * 1024 * 1024
_RISKY = ("=", "+", "-", "@", "\t", "\r")


class InvalidTable(AppError):
    status_code, code = 422, "invalid_file"


def _clean(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    if len(text) > MAX_CELL:
        raise InvalidTable("cell_too_long")
    return text


def _read_csv(raw: bytes) -> list[list[str]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise InvalidTable("not_utf8") from None
    sample = text[:4096]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","  # Excel in ID uses ;
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(raw: bytes) -> list[list[Any]]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED:
                raise InvalidTable("file_too_large")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError):
        raise InvalidTable() from None
    try:
        sheet = book.worksheets[0]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_ROWS + 1:
                break
        return rows
    finally:
        book.close()


def _record(header: list[str], row: Sequence[Any]) -> dict[str, str]:
    return {h: _clean(row[i]) if i < len(row) else "" for i, h in enumerate(header) if h}


def read_table(raw: bytes, file_name: str) -> list[dict[str, str]]:
    """First row is the header (lower-cased); returns one dict per non-empty row."""
    rows = _read_xlsx(raw) if file_name.lower().endswith(".xlsx") else _read_csv(raw)
    if not rows:
        raise InvalidTable("empty")
    header = [_clean(h).lower() for h in rows[0]]
    if len(header) > MAX_COLUMNS:
        raise InvalidTable("too_many_columns")
    records = [_record(header, r) for r in rows[1:]]
    body = [r for r in records if any(r.values())]
    if len(body) > MAX_ROWS:
        raise InvalidTable("too_many_rows", details={"max_rows": MAX_ROWS})
    return body


def _safe(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(_RISKY):
        return "'" + value
    return value


def write_csv(header: Sequence[str], rows: Iterable[Sequence[Any]]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(header)
    writer.writerows([_safe(v) for v in row] for row in rows)
    return out.getvalue().encode("utf-8-sig")  # BOM: Excel opens UTF-8 correctly


def write_xlsx(header: Sequence[str], rows: Iterable[Sequence[Any]]) -> bytes:
    book = Workbook(write_only=True)
    sheet = book.create_sheet()
    sheet.append(list(header))
    for row in rows:
        sheet.append([_safe(v) for v in row])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
